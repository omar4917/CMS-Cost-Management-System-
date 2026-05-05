from __future__ import annotations

import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT_DIR = Path(__file__).resolve().parents[1]
WORKBOOK_PATH = ROOT_DIR.parents[1] / "Cost Tracker_V1.1.xlsx"
MANIFEST_PATH = ROOT_DIR / "data" / "cost_tracker_v1_1_import_manifest.json"
SOURCE_TABLE = "tracker_import_cost_tracker_v1_1"
SOURCE_MARKER = "Imported from Cost Tracker_V1.1.xlsx"

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from core.database import ensure_cash_transactions_table, ensure_project_profile_columns, get_connection


TARGET_PROJECT_NAME = "RRDL Office Cost"

PROJECT_META = {
    TARGET_PROJECT_NAME: {
        "type": "commercial",
        "status": "active",
        "property_type": "Office Overhead Ledger",
        "property_for": "Other",
        "construction_status": "Operational",
        "transaction_type": "New",
        "description_prefix": "Single CMS project imported from the source cost tracker workbook.",
    }
}


MONTH_RE = re.compile(r"^(?P<month>[A-Za-z]+)_(?P<year>\d{2})$")
PROJECT_BUCKET_TOKEN_RE = re.compile(r"(?<!\w)L-\d+(?!\w)", re.IGNORECASE)
PAYMENT_METHOD_RULES = [
    ("mobile_banking", ("bkash", "nagad", "rocket", "mobile")),
    ("cheque", ("cheque", "check", "a/c payee")),
    ("bank_transfer", ("transfer", "bank", "deposit", "atm", "online")),
    ("cash", ("cash", "hand cash", "paid by own", "payment")),
]


def clean_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).replace("\r", "\n").strip()
    if not text or text.lower() == "nan":
        return None
    return text


def strip_bucket_label(value, bucket_label: str | None) -> str | None:
    text = clean_text(value)
    if not text:
        return None

    cleaned = text
    label = clean_text(bucket_label)
    if label and label.lower() != TARGET_PROJECT_NAME.lower():
        cleaned = re.sub(rf"(?<!\w){re.escape(label)}(?!\w)", "", cleaned, flags=re.IGNORECASE)
    cleaned = PROJECT_BUCKET_TOKEN_RE.sub("", cleaned)

    cleaned = re.sub(r"[ \t]{2,}", " ", cleaned)
    cleaned = re.sub(r"\s+\n", "\n", cleaned)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned)
    cleaned = re.sub(r"\(\s*\)", "", cleaned)
    cleaned = re.sub(r"\[\s*\]", "", cleaned)
    cleaned = re.sub(r"\{\s*\}", "", cleaned)
    cleaned = re.sub(r"\s+([,;:|/\-)])", r"\1", cleaned)
    cleaned = re.sub(r"([(/-])\s+", r"\1", cleaned)
    cleaned = re.sub(r"^[,;:|/\-\s]+", "", cleaned)
    cleaned = re.sub(r"[,;:|/\-\s]+$", "", cleaned)
    return clean_text(cleaned)


def clean_amount(value) -> float:
    if value in (None, ""):
        return 0.0
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return 0.0


def clean_int_like(value) -> float:
    number = clean_amount(value)
    if abs(number) < 1e-9:
        return 0.0
    return number


def parse_tracker_month(month_tag: str | None) -> tuple[int, int] | None:
    tag = clean_text(month_tag)
    if not tag:
        return None
    match = MONTH_RE.match(tag)
    if not match:
        return None

    month_text = match.group("month")
    year = 2000 + int(match.group("year"))
    for fmt in ("%b", "%B"):
        try:
            month = datetime.strptime(month_text, fmt).month
            return year, month
        except ValueError:
            continue
    return None


def to_datetime(value) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value

    text = clean_text(value)
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def normalize_date(value, month_tag: str | None) -> tuple[str | None, bool]:
    dt = to_datetime(value)
    if not dt:
        return None, False

    corrected = False
    tracker_month = parse_tracker_month(month_tag)
    if tracker_month:
        tracker_year, tracker_month_num = tracker_month
        if dt.year == tracker_year + 1 and dt.month == tracker_month_num:
            dt = dt.replace(year=tracker_year)
            corrected = True

    return dt.date().isoformat(), corrected


def infer_payment_method(receive_detail: str | None, received_from: str | None) -> str | None:
    haystack = " ".join(part for part in [receive_detail or "", received_from or ""] if part).lower()
    for method, needles in PAYMENT_METHOD_RULES:
        if any(needle in haystack for needle in needles):
            return method
    return None


def build_project_description(name: str, summary: dict[str, object]) -> str:
    meta = PROJECT_META.get(name, {})
    prefix = meta.get("description_prefix", "Imported project bucket from the source cost tracker.")
    lines = [
        prefix,
        f"Source workbook: {WORKBOOK_PATH.name}.",
        f"Imported tracker rows: {summary['rows']}.",
        f"Recorded receive rows: {summary['receive_rows']} (BDT {summary['receive_total']:,.2f}).",
        f"Recorded cost rows: {summary['cost_rows']} (BDT {summary['cost_total']:,.2f}).",
        "Progress was left at 0 because the workbook does not contain an explicit completion percentage.",
    ]
    return " ".join(lines)


def load_workbook_rows() -> tuple[list[dict[str, object]], dict[str, dict[str, object]], list[int]]:
    workbook = load_workbook(WORKBOOK_PATH, data_only=True)
    sheet = workbook["Details Cost"]

    summaries: dict[str, dict[str, object]] = defaultdict(
        lambda: {
            "rows": 0,
            "receive_rows": 0,
            "cost_rows": 0,
            "receive_total": 0.0,
            "cost_total": 0.0,
            "start_date": None,
            "end_date": None,
        }
    )
    rows: list[dict[str, object]] = []
    corrected_rows: list[int] = []

    for row_index in range(4, sheet.max_row + 1):
        values = [sheet.cell(row=row_index, column=column).value for column in range(1, 21)]
        (
            receive_date_raw,
            receive_detail_raw,
            receive_amount_raw,
            received_from_raw,
            cost_date_raw,
            cost_detail_raw,
            cost_amount_raw,
            pay_to_raw,
            unit_raw,
            unit_rate_raw,
            qty_raw,
            qty_cft_raw,
            remarks_raw,
            cost_head_materials_raw,
            structure_raw,
            cost_summary_raw,
            boq_raw,
            floors_raw,
            project_bucket_raw,
            month_raw,
        ) = values

        source_bucket = clean_text(project_bucket_raw)
        receive_detail = strip_bucket_label(receive_detail_raw, source_bucket)
        receive_amount = clean_amount(receive_amount_raw)
        received_from = strip_bucket_label(received_from_raw, source_bucket)
        cost_detail = strip_bucket_label(cost_detail_raw, source_bucket)
        cost_amount = clean_amount(cost_amount_raw)
        pay_to = strip_bucket_label(pay_to_raw, source_bucket)
        month_tag = clean_text(month_raw)
        project_bucket = strip_bucket_label(project_bucket_raw, source_bucket)

        if not any(
            [
                receive_detail,
                abs(receive_amount) > 0,
                received_from,
                cost_detail,
                abs(cost_amount) > 0,
                pay_to,
                project_bucket,
            ]
        ):
            continue

        receive_date, receive_corrected = normalize_date(receive_date_raw, month_tag)
        cost_date, cost_corrected = normalize_date(cost_date_raw, month_tag)
        if receive_corrected or cost_corrected:
            corrected_rows.append(row_index)

        row = {
            "row_index": row_index,
            "project_bucket": project_bucket,
            "receive_date": receive_date,
            "receive_detail": receive_detail,
            "receive_amount": receive_amount,
            "received_from": received_from,
            "cost_date": cost_date,
            "cost_detail": cost_detail,
            "cost_amount": cost_amount,
            "pay_to": pay_to,
            "unit": strip_bucket_label(unit_raw, source_bucket),
            "unit_rate": clean_amount(unit_rate_raw),
            "qty": clean_int_like(qty_raw),
            "qty_cft": clean_int_like(qty_cft_raw),
            "remarks": strip_bucket_label(remarks_raw, source_bucket),
            "cost_head_materials": strip_bucket_label(cost_head_materials_raw, source_bucket),
            "structure_or_finishing": strip_bucket_label(structure_raw, source_bucket),
            "cost_summary_1": strip_bucket_label(cost_summary_raw, source_bucket),
            "category_boq_mapping": strip_bucket_label(boq_raw, source_bucket),
            "cost_head_floors": strip_bucket_label(floors_raw, source_bucket),
            "cost_head_project": TARGET_PROJECT_NAME,
            "cost_head_months": month_tag,
            "payment_method": infer_payment_method(receive_detail, received_from),
        }
        rows.append(row)

        summary = summaries[TARGET_PROJECT_NAME]
        summary["rows"] += 1
        if abs(receive_amount) > 0:
            summary["receive_rows"] += 1
            summary["receive_total"] += receive_amount
        if cost_detail or abs(cost_amount) > 0:
            summary["cost_rows"] += 1
            summary["cost_total"] += cost_amount

        candidate = cost_date or receive_date
        if candidate:
            if summary["start_date"] is None or candidate < summary["start_date"]:
                summary["start_date"] = candidate
            if summary["end_date"] is None or candidate > summary["end_date"]:
                summary["end_date"] = candidate

    return rows, summaries, corrected_rows


def fetch_created_by(cursor) -> int | None:
    cursor.execute(
        """
        SELECT id
        FROM users
        WHERE LOWER(COALESCE(username, '')) = 'admin'
        ORDER BY id ASC
        LIMIT 1
        """
    )
    row = cursor.fetchone()
    if row:
        return row["id"]

    cursor.execute("SELECT id FROM users ORDER BY id ASC LIMIT 1")
    row = cursor.fetchone()
    return row["id"] if row else None


def load_manifest() -> dict[str, object]:
    if not MANIFEST_PATH.exists():
        return {}
    try:
        return json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def delete_ids(cursor, table_name: str, ids: list[int]) -> None:
    clean_ids = [int(value) for value in ids if value]
    if not clean_ids:
        return
    placeholders = ",".join("?" for _ in clean_ids)
    cursor.execute(f"DELETE FROM `{table_name}` WHERE `id` IN ({placeholders})", clean_ids)


def upsert_project(cursor, created_by: int | None, name: str, summary: dict[str, object]) -> int:
    meta = PROJECT_META.get(name, {})
    payload = {
        "description": build_project_description(name, summary),
        "location": None,
        "type": meta.get("type", "mixed"),
        "property_type": meta.get("property_type"),
        "property_for": meta.get("property_for"),
        "status": meta.get("status", "active"),
        "construction_status": meta.get("construction_status"),
        "start_date": summary.get("start_date"),
        "estimated_end_date": summary.get("end_date"),
        "actual_end_date": None,
        "total_budget": round(float(summary.get("cost_total") or 0), 2),
        "progress": 0,
        "unit_size": None,
        "transaction_type": meta.get("transaction_type"),
        "floor_available_on": None,
        "bedrooms": None,
        "bathrooms": None,
        "balconies": None,
        "garages": None,
        "total_floors": None,
        "total_units": None,
        "furnishing": None,
        "facing": None,
        "land_area": None,
        "building_area": None,
        "features": None,
        "nearby_places": None,
        "corporate_office": None,
        "cover_image_path": None,
    }

    cursor.execute("SELECT id FROM projects WHERE name = ? LIMIT 1", (name,))
    existing = cursor.fetchone()
    if existing:
        project_id = int(existing["id"])
        assignments = ", ".join(f"`{column}` = ?" for column in payload)
        cursor.execute(
            f"UPDATE projects SET {assignments}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
            tuple(payload.values()) + (project_id,),
        )
        return project_id

    columns = ["name", *payload.keys(), "created_by", "created_at", "updated_at"]
    values = [name, *payload.values(), created_by, "CURRENT_TIMESTAMP", "CURRENT_TIMESTAMP"]
    placeholders = ["?"] * (len(columns) - 2) + ["CURRENT_TIMESTAMP", "CURRENT_TIMESTAMP"]
    cursor.execute(
        f"INSERT INTO projects ({', '.join(columns)}) VALUES ({', '.join(placeholders)})",
        tuple(values[:-2]),
    )
    return int(cursor.lastrowid)


def insert_cost_item(cursor, created_by: int | None, project_id: int | None, row: dict[str, object]) -> int:
    source_note = f"{SOURCE_MARKER}; workbook row {row['row_index']}"
    name = row["cost_detail"] or row["receive_detail"] or f"Imported Row {row['row_index']}"
    description = row["cost_detail"] or row["receive_detail"]
    generic_date = row["cost_date"] or row["receive_date"]
    vendor = row["pay_to"] or row["received_from"]
    status = "approved" if abs(float(row["cost_amount"] or 0)) > 0 else "pending"

    cursor.execute(
        """
        INSERT INTO cost_items (
            name, project_id, category_id, description, quantity, unit, unit_price,
            estimated_amount, actual_amount, currency, date, vendor, invoice_no, status,
            notes, created_by,
            receive_date, receive_detail, receive_amount, received_from,
            cost_date, cost_detail, cost_amount, pay_to, unit_rate, qty, qty_cft, remarks,
            cost_head_materials, structure_or_finishing, cost_summary_1, category_boq_mapping,
            cost_head_floors, cost_head_project, cost_head_months, created_at, updated_at
        ) VALUES (
            ?, ?, NULL, ?, ?, ?, ?,
            ?, ?, 'BDT', ?, ?, ?, ?,
            ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, ?, ?, ?, ?, ?,
            ?, ?, ?, ?,
            ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP
        )
        """,
        (
            name,
            project_id,
            description,
            row["qty"],
            row["unit"],
            row["unit_rate"],
            row["cost_amount"],
            row["cost_amount"],
            generic_date,
            vendor,
            f"xlsx-row-{row['row_index']}",
            status,
            source_note,
            created_by,
            row["receive_date"],
            row["receive_detail"],
            row["receive_amount"],
            row["received_from"],
            row["cost_date"],
            row["cost_detail"],
            row["cost_amount"],
            row["pay_to"],
            row["unit_rate"],
            row["qty"],
            row["qty_cft"],
            row["remarks"],
            row["cost_head_materials"],
            row["structure_or_finishing"],
            row["cost_summary_1"],
            row["category_boq_mapping"],
            row["cost_head_floors"],
            row["cost_head_project"],
            row["cost_head_months"],
        ),
    )
    return int(cursor.lastrowid)


def insert_receive_transaction(cursor, created_by: int | None, project_id: int | None, row: dict[str, object]) -> int | None:
    receive_amount = float(row["receive_amount"] or 0)
    if abs(receive_amount) < 1e-9:
        return None

    direction = "inflow" if receive_amount >= 0 else "outflow"
    entry_type = "other_inflow" if direction == "inflow" else "refund_out"
    counterparty = row["received_from"] or row["receive_detail"] or TARGET_PROJECT_NAME
    notes = f"{SOURCE_MARKER}; workbook row {row['row_index']}"
    if row["receive_detail"]:
        notes = f"{row['receive_detail']} | {notes}"

    cursor.execute(
        """
        INSERT INTO cash_transactions (
            project_id, investor_id, contractor_id, entry_type, direction, amount, currency,
            tx_date, counterparty, reference_no, payment_method, status, source_table,
            source_id, notes, created_by, created_at, updated_at
        ) VALUES (
            ?, NULL, NULL, ?, ?, ?, 'BDT',
            ?, ?, ?, ?, 'confirmed', ?, NULL, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP
        )
        """,
        (
            project_id,
            entry_type,
            direction,
            abs(receive_amount),
            row["receive_date"] or row["cost_date"],
            counterparty,
            f"CT-V1.1-R{row['row_index']}",
            row["payment_method"],
            SOURCE_TABLE,
            notes,
            created_by,
        ),
    )
    return int(cursor.lastrowid)


def main() -> None:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(f"Workbook not found: {WORKBOOK_PATH}")

    conn = get_connection()
    ensure_project_profile_columns()
    ensure_cash_transactions_table()

    rows, summaries, corrected_rows = load_workbook_rows()
    if not rows:
        raise RuntimeError("No workbook rows were parsed.")

    cursor = conn.cursor()

    manifest = load_manifest()

    try:
        delete_ids(cursor, "cash_transactions", manifest.get("cash_transaction_ids", []))
        delete_ids(cursor, "cost_items", manifest.get("cost_item_ids", []))

        created_by = fetch_created_by(cursor)
        project_ids: dict[str, int] = {}
        project_ids[TARGET_PROJECT_NAME] = upsert_project(
            cursor,
            created_by,
            TARGET_PROJECT_NAME,
            summaries[TARGET_PROJECT_NAME],
        )

        cost_item_ids: list[int] = []
        cash_transaction_ids: list[int] = []
        imported_rows = 0
        receive_rows = 0

        for row in rows:
            project_id = project_ids[TARGET_PROJECT_NAME]
            cost_item_ids.append(insert_cost_item(cursor, created_by, project_id, row))
            imported_rows += 1

            cash_id = insert_receive_transaction(cursor, created_by, project_id, row)
            if cash_id:
                cash_transaction_ids.append(cash_id)
                receive_rows += 1

        conn.commit()

        manifest_payload = {
            "source_workbook": str(WORKBOOK_PATH),
            "source_table": SOURCE_TABLE,
            "imported_at": datetime.now().isoformat(timespec="seconds"),
            "project_ids": project_ids,
            "cost_item_ids": cost_item_ids,
            "cash_transaction_ids": cash_transaction_ids,
            "corrected_date_rows": corrected_rows,
            "summary": summaries,
        }
        MANIFEST_PATH.parent.mkdir(parents=True, exist_ok=True)
        MANIFEST_PATH.write_text(json.dumps(manifest_payload, indent=2, ensure_ascii=True), encoding="utf-8")

        print(f"Imported workbook: {WORKBOOK_PATH.name}")
        print(f"Projects upserted: {len(project_ids)}")
        print(f"Cost tracker rows imported: {imported_rows}")
        print(f"Receive-side cash transactions imported: {receive_rows}")
        print(f"Date corrections applied: {len(corrected_rows)}")
        print(f"Manifest written: {MANIFEST_PATH}")
        for project_name, project_id in sorted(project_ids.items(), key=lambda item: item[1]):
            summary = summaries[project_name]
            print(
                f"- Project #{project_id} {project_name}: "
                f"{summary['rows']} rows | "
                f"receive BDT {summary['receive_total']:,.2f} | "
                f"cost BDT {summary['cost_total']:,.2f}"
            )
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()


if __name__ == "__main__":
    main()
